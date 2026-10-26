"""Regenerates the synthetic test emails in this folder.

Every name, address and figure here is fictional, so the fixtures are safe to
commit (CLAUDE.md rule 2). Real samples never go into tests/fixtures as-is:
rewrite them into this script with invented names and amounts instead.

    uv run python tests/fixtures/make_fixtures.py
"""
from __future__ import annotations

import io
from email import policy
from email.message import EmailMessage
from pathlib import Path

HERE = Path(__file__).parent
XLSX = ("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet")
DOCX = ("application", "vnd.openxmlformats-officedocument.wordprocessingml.document")


def email_message(sender: str, subject: str, body: str, html: str | None = None,
                  date: str = "Mon, 21 Sep 2026 09:12:00 +0200") -> EmailMessage:
    msg = EmailMessage()
    msg["From"] = sender
    msg["To"] = "Actuarial Team <actuarial@example.com>"
    msg["Date"] = date
    msg["Subject"] = subject
    msg.set_content(body)
    if html:
        msg.add_alternative(html, subtype="html")
    return msg


def xlsx(rows: list[list]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def docx(paragraphs: list[str], table: list[list[str]]) -> bytes:
    import docx as python_docx

    doc = python_docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    t = doc.add_table(rows=0, cols=len(table[0]))
    for row in table:
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pdf(lines: list[str]) -> bytes:
    """A minimal one-page PDF with a real text layer."""
    esc = lambda s: s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = "BT /F1 11 Tf 50 780 Td 14 TL " + " ".join(f"({esc(line)}) Tj T*" for line in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def smoke_test() -> EmailMessage:
    msg = email_message("Pipeline Tester <tester@example.com>", "[PIPELINE TEST] hello",
                        "Hello pipeline,\n\nthis is an end-to-end test with one small attachment. Please ignore.\n")
    msg.add_attachment(b"item,value\napples,3\npears,5\n", maintype="text", subtype="csv", filename="test_numbers.csv")
    return msg


def large_loss_advice() -> EmailMessage:
    body = """Dear colleagues,

On behalf of Nordwind Mutual Insurance (fictional) we advise the following large loss
under the Property Catastrophe Excess of Loss programme 2026, Layer 2
(treaty ref. NWM-PCAT-2026-L2, underwriting year 2026).

Event: Storm Example (windstorm)
Date of loss: 14 September 2026
Country: Exampleland
Line of business: Property
Current gross loss estimate (100%): EUR 12.5m
Previous estimate: EUR 10,000,000
Your share: 7.5%

The cedent's loss advice and the loss details are attached. Please confirm receipt.

Kind regards
Claims Desk, Example Broker Ltd (fictional)
"""
    html = """<html><head><style>p {margin:0}</style></head><body>
<p>Dear colleagues,</p>
<p>On behalf of <b>Nordwind Mutual Insurance</b> (fictional) we advise the following large loss
under the Property Catastrophe Excess of Loss programme 2026, Layer 2
(treaty ref. NWM-PCAT-2026-L2, underwriting year 2026).</p>
<table border="1">
<tr><th>Item</th><th>Value</th></tr>
<tr><td>Event</td><td>Storm Example (windstorm)</td></tr>
<tr><td>Date of loss</td><td>14 September 2026</td></tr>
<tr><td>Country</td><td>Exampleland</td></tr>
<tr><td>Line of business</td><td>Property</td></tr>
<tr><td>Current gross loss estimate (100%)</td><td>EUR 12.5m</td></tr>
<tr><td>Previous estimate</td><td>EUR 10,000,000</td></tr>
<tr><td>Your share</td><td>7.5%</td></tr>
</table>
<p>The cedent's loss advice and the loss details are attached. Please confirm receipt.</p>
<p>Kind regards<br>Claims Desk, Example Broker Ltd (fictional)</p>
</body></html>"""
    msg = email_message("Claims Desk <claims@examplebroker.test>",
                        "Large Loss Advice - Nordwind Mutual - DOL 14/09/2026", body, html)
    msg.add_attachment(pdf([
        "LOSS ADVICE - Nordwind Mutual Insurance (fictional)",
        "Treaty: NWM-PCAT-2026-L2   Underwriting year: 2026",
        "Event: Storm Example   Date of loss: 14.09.2026",
        "Gross loss estimate (100%): EUR 12,500,000",
        "Previous estimate: EUR 10,000,000",
    ]), maintype="application", subtype="pdf", filename="loss_advice.pdf")
    msg.add_attachment(xlsx([
        ["Policy", "Location", "Gross loss EUR"],
        ["P-0001", "Harbour Town", 7_250_000],
        ["P-0002", "Hill Village", 5_250_000],
    ]), maintype=XLSX[0], subtype=XLSX[1], filename="loss_details.xlsx")
    return msg


def loss_update_forwarded() -> EmailMessage:
    inner = email_message(
        "Nordwind Mutual Reinsurance Dept <re@nordwind.test>",
        "Loss update: Storm Example - revised estimate",
        "Dear reinsurers,\n\nfor treaty NWM-PCAT-2026-L2 (UY 2026), loss date 14.09.2026, Storm Example,\n"
        "we revise our gross loss estimate to EUR 15.0m (previously EUR 12.5m).\n"
        "A detailed report will follow.\n\nNordwind Mutual Insurance (fictional)\n",
        date="Fri, 2 Oct 2026 16:40:00 +0200",
    )
    outer = email_message("A. Colleague <colleague@example.com>", "FW: Loss update Storm Example",
                          "FYI, see the attached update from the cedent.\n", date="Mon, 5 Oct 2026 08:05:00 +0200")
    outer.add_attachment(inner)
    return outer


def cash_call() -> EmailMessage:
    msg = email_message("Nordwind Mutual Reinsurance Dept <re@nordwind.test>",
                        "Cash call request - Nordwind Mutual - Storm Example",
                        "Dear reinsurers,\n\nplease find attached our cash call request for Storm Example.\n\n"
                        "Nordwind Mutual Insurance (fictional)\n", date="Mon, 12 Oct 2026 10:00:00 +0200")
    msg.add_attachment(docx(
        ["CASH CALL REQUEST", "Nordwind Mutual Insurance (fictional) requests the following cash call."],
        [["Treaty", "NWM-PCAT-2026-L2"], ["Event", "Storm Example, 14 September 2026"],
         ["Your share", "7.5%"], ["Cash call amount", "EUR 375,000"], ["Due date", "30 October 2026"]],
    ), maintype=DOCX[0], subtype=DOCX[1], filename="cash_call_letter.docx")
    return msg


def quarterly_accounts() -> EmailMessage:
    msg = email_message("Accounts <accounts@examplelife.test>",
                        "Q3 2026 technical accounts - Example Life Re Treaty (fictional)",
                        "Dear partners,\n\nplease find attached the Q3 2026 technical account statement for the "
                        "quota share treaty.\nBalance in your favour: EUR 48,210.55.\n\nExample Life (fictional)\n",
                        date="Thu, 15 Oct 2026 11:30:00 +0200")
    msg.add_attachment(xlsx([
        ["Item", "EUR"], ["Premium", 412_000.00], ["Commission", -103_000.00],
        ["Claims paid", -260_789.45], ["Balance", 48_210.55],
    ]), maintype=XLSX[0], subtype=XLSX[1], filename="Q3_2026_statement.xlsx")
    return msg


FIXTURES = {
    "smoke_test.eml": smoke_test,
    "large_loss_advice.eml": large_loss_advice,
    "loss_update_forwarded.eml": loss_update_forwarded,
    "cash_call.eml": cash_call,
    "quarterly_accounts.eml": quarterly_accounts,
}

if __name__ == "__main__":
    for name, build in FIXTURES.items():
        (HERE / name).write_bytes(build().as_bytes(policy=policy.SMTP))
        print("wrote", name)
